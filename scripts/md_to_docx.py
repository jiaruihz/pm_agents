#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError as exc:
    raise SystemExit(
        "Missing dependency: python-docx\n"
        "Install it in a virtual environment, for example:\n"
        "  python3 -m venv .venv-md-docx\n"
        "  .venv-md-docx/bin/pip install python-docx"
    ) from exc


HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
UNORDERED_RE = re.compile(r"^(\s*)[-+*]\s+(.*)$")
ORDERED_RE = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
BLOCKQUOTE_RE = re.compile(r"^>\s?(.*)$")
HR_RE = re.compile(r"^\s*([-*_])(?:\s*\1){2,}\s*$")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def set_east_asia_font(run, font_name: str) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for style_name in ("Quote", "Intense Quote"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(11)
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")

    if "CodeBlock" not in doc.styles:
        code_style = doc.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
        code_style.base_style = doc.styles["Normal"]
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(10.5)
        code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "等宽更纱黑体 SC")


def add_runs(paragraph, text: str, default_font: str = "宋体") -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_east_asia_font(run, default_font)

        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10.5)
            run.font.highlight_color = None
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True

        set_east_asia_font(run, default_font)
        cursor = match.end()

    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_east_asia_font(run, default_font)


def add_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    add_runs(paragraph, text)


def add_heading(doc: Document, level: int, text: str) -> None:
    if level == 1:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        set_east_asia_font(run, "黑体")
        paragraph.paragraph_format.space_after = Pt(12)
        paragraph.paragraph_format.space_before = Pt(0)
        return

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(14 if level == 2 else 12)
    set_east_asia_font(run, "黑体")


def add_list_item(doc: Document, text: str, ordered: bool, indent_spaces: int) -> None:
    style_name = "List Number" if ordered else "List Bullet"
    paragraph = doc.add_paragraph(style=style_name)
    paragraph.paragraph_format.left_indent = Cm(0.74 + (indent_spaces // 2) * 0.74)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    add_runs(paragraph, text)


def add_quote(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Quote" if "Quote" in doc.styles else None)
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.right_indent = Cm(0.74)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    add_runs(paragraph, text, default_font="楷体")


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines or [""]:
        paragraph = doc.add_paragraph(style="CodeBlock")
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.right_indent = Cm(0.74)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(10.5)
        set_east_asia_font(run, "等宽更纱黑体 SC")


def convert_markdown(markdown_text: str) -> Document:
    doc = Document()
    configure_document(doc)

    in_code_block = False
    code_lines: list[str] = []

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()

        if line.startswith("```"):
            if in_code_block:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if not line.strip():
            continue

        if HR_RE.match(line):
            continue

        if match := HEADING_RE.match(line):
            add_heading(doc, len(match.group(1)), match.group(2).strip())
            continue

        if match := UNORDERED_RE.match(line):
            add_list_item(doc, match.group(2).strip(), ordered=False, indent_spaces=len(match.group(1)))
            continue

        if match := ORDERED_RE.match(line):
            add_list_item(doc, match.group(2).strip(), ordered=True, indent_spaces=len(match.group(1)))
            continue

        if match := BLOCKQUOTE_RE.match(line):
            add_quote(doc, match.group(1).strip())
            continue

        add_paragraph(doc, line.strip())

    if in_code_block:
        add_code_block(doc, code_lines)

    return doc


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert a Markdown file to a readable DOCX.")
    parser.add_argument("input", type=Path, help="Input Markdown file")
    parser.add_argument("output", nargs="?", type=Path, help="Output DOCX file")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path: Path = args.input
    output_path: Path = args.output or input_path.with_suffix(".docx")

    if not input_path.is_file():
        print(f"Input file not found: {input_path}", file=sys.stderr)
        return 1

    markdown_text = input_path.read_text(encoding="utf-8")
    document = convert_markdown(markdown_text)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
